import os
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

DOCX_OUT = "/home/malulu/proj/enhanced-sw/report/final/enhanced_sw_project_report_concise.docx"
REPORT_MD = "/home/malulu/proj/enhanced-sw/report/draft/report_concise.md"
BASE_DIR = "/home/malulu/proj/enhanced-sw"

os.makedirs("/home/malulu/proj/enhanced-sw/report/final", exist_ok=True)

def set_cell_background(cell, hex_color):
    shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{hex_color}"/>')
    cell._tc.get_or_add_tcPr().append(shading_elm)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def clean_math_text(text):
    res = text
    res = res.replace('<br>', '\n').replace('<br/>', '\n')
    res = re.sub(r'\\frac\{([^}]+)\}\{([^}]+)\}', r'(\1 / \2)', res)
    res = re.sub(r'\\left\(', '(', res).replace(r'\right)', ')')
    res = re.sub(r'\\sqrt\{([^}]+)\}', r'sqrt(\1)', res)
    res = res.replace('\\Delta', 'Δ').replace('\\neq', '≠').replace('\\ne', '≠').replace('\\equiv', '≡')
    res = res.replace('\\times', '×').replace('\\cdot', ' · ').replace('\\log', 'log').replace('\\implies', ' ⟹ ')
    res = res.replace('\\le', '≤').replace('\\ge', '≥').replace('\\leq', '≤').replace('\\geq', '≥')
    res = res.replace('\\sum', '∑').replace('\\mu', 'μ').replace('\\sigma', 'σ').replace('\\epsilon', 'ε')
    res = res.replace('\\in', '∈').replace('\\to', '→').replace('\\hat{y}', 'ŷ')
    res = re.sub(r'\\text\{([^}]+)\}', r'\1', res)
    res = re.sub(r'\\mathrm\{([^}]+)\}', r'\1', res)
    res = res.replace('\\mathbb{Q}^+', 'ℚ⁺').replace('\\mathbb{R}', 'ℝ').replace('\\dots', '…')
    res = res.replace('\\max', 'max').replace('\\min', 'min')
    res = res.replace('\\%', '%').replace('\\{', '{').replace('\\}', '}')
    res = re.sub(r'_\{([^}]+)\}', r'_\1', res)
    res = res.replace('$', '')
    return res

def build_docx():
    doc = Document()
    
    # Standard A4 with 0.8-inch margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)
        section.page_width = Inches(8.27)
        section.page_height = Inches(11.69)
        
        # Header left empty - no running header
        header = section.header
        p_hdr = header.paragraphs[0]
        p_hdr.text = ""

    # Base Styles
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(10)
    normal_style.font.color.rgb = RGBColor(30, 41, 59)

    with open(REPORT_MD, 'r', encoding='utf-8') as f:
        md_text = f.read()

    lines = md_text.split('\n')
    i = 0
    in_code_block = False
    code_lines = []

    while i < len(lines):
        line = lines[i]

        # Handle Code Blocks
        if line.startswith('```'):
            if in_code_block:
                table = doc.add_table(rows=1, cols=1)
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                cell = table.cell(0, 0)
                set_cell_background(cell, "F8FAFC")
                set_cell_margins(cell, top=100, bottom=100, left=150, right=150)
                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(4)
                run = p.add_run("\n".join(code_lines))
                run.font.name = 'Consolas'
                run.font.size = Pt(8.5)
                run.font.color.rgb = RGBColor(15, 23, 42)
                doc.add_paragraph().paragraph_format.space_after = Pt(6)
                in_code_block = False
                code_lines = []
            else:
                in_code_block = True
                code_lines = []
            i += 1
            continue

        if in_code_block:
            code_lines.append(line)
            i += 1
            continue

        # Handle Tables
        if line.strip().startswith('|') and line.count('|') >= 2:
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith('|') and lines[i].count('|') >= 2:
                table_lines.append(lines[i].strip())
                i += 1
            
            if len(table_lines) >= 3:
                raw_headers = [c.strip() for c in table_lines[0].split('|')[1:-1]]
                raw_rows = []
                for tl in table_lines[2:]:
                    cells = [c.strip() for c in tl.split('|')[1:-1]]
                    if cells:
                        raw_rows.append(cells)
                
                table = doc.add_table(rows=len(raw_rows) + 1, cols=len(raw_headers))
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                
                # Format Header
                hdr_cells = table.rows[0].cells
                for idx, text in enumerate(raw_headers):
                    hdr_cells[idx].text = clean_math_text(text)
                    set_cell_background(hdr_cells[idx], "1E293B")
                    set_cell_margins(hdr_cells[idx], top=80, bottom=80, left=100, right=100)
                    p = hdr_cells[idx].paragraphs[0]
                    p.paragraph_format.space_before = Pt(2)
                    p.paragraph_format.space_after = Pt(2)
                    for run in p.runs:
                        run.font.bold = True
                        run.font.size = Pt(8)
                        run.font.color.rgb = RGBColor(255, 255, 255)

                # Format Rows
                for r_idx, row_data in enumerate(raw_rows):
                    row_cells = table.rows[r_idx + 1].cells
                    bg_color = "FFFFFF" if r_idx % 2 == 0 else "F8FAFC"
                    for c_idx, cell_value in enumerate(row_data):
                        if c_idx < len(row_cells):
                            row_cells[c_idx].text = clean_math_text(cell_value)
                            set_cell_background(row_cells[c_idx], bg_color)
                            set_cell_margins(row_cells[c_idx], top=60, bottom=60, left=100, right=100)
                            p = row_cells[c_idx].paragraphs[0]
                            p.paragraph_format.space_before = Pt(1)
                            p.paragraph_format.space_after = Pt(1)
                            for run in p.runs:
                                run.font.size = Pt(7.5)
                                run.font.color.rgb = RGBColor(51, 65, 85)

                doc.add_paragraph().paragraph_format.space_after = Pt(6)
            continue

        # Handle Images: ![caption](path)
        img_match = re.match(r'^!\[(.*?)\]\((.*?)\)', line.strip())
        if img_match:
            caption = img_match.group(1)
            img_rel_path = img_match.group(2)
            img_full_path = os.path.join(BASE_DIR, img_rel_path) if not os.path.isabs(img_rel_path) else img_rel_path
            
            explicit_caption = caption
            if i + 1 < len(lines) and lines[i + 1].strip().startswith('*') and lines[i + 1].strip().endswith('*'):
                explicit_caption = lines[i + 1].strip().strip('*')
                i += 1

            if os.path.exists(img_full_path):
                p_img = doc.add_paragraph()
                p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_img.paragraph_format.space_before = Pt(8)
                p_img.paragraph_format.space_after = Pt(2)
                run = p_img.add_run()
                run.add_picture(img_full_path, width=Inches(5.4))
                
                # Caption
                p_cap = doc.add_paragraph()
                p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_cap.paragraph_format.space_after = Pt(8)
                run_cap = p_cap.add_run(explicit_caption)
                run_cap.font.size = Pt(8.5)
                run_cap.font.italic = True
                run_cap.font.color.rgb = RGBColor(100, 116, 139)
            i += 1
            continue

        # Check Headings
        if line.startswith('# '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(16)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(line[2:].strip())
            run.font.bold = True
            run.font.size = Pt(16)
            run.font.color.rgb = RGBColor(15, 23, 42)
        elif line.startswith('## '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(line[3:].strip())
            run.font.bold = True
            run.font.size = Pt(13)
            run.font.color.rgb = RGBColor(30, 58, 138)
        elif line.startswith('### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(2)
            run = p.add_run(line[4:].strip())
            run.font.bold = True
            run.font.size = Pt(11)
            run.font.color.rgb = RGBColor(51, 65, 85)
        elif line.startswith('---'):
            pass
        elif line.strip() == '':
            pass
        else:
            # Paragraph or list
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.15
            
            clean_line = clean_math_text(line.strip())
            parts = re.split(r'(\*\*.*?\*\*|\*.*?\*|`.*?`)', clean_line)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    run = p.add_run(part[2:-2])
                    run.font.bold = True
                elif part.startswith('*') and part.endswith('*'):
                    run = p.add_run(part[1:-1])
                    run.font.italic = True
                elif part.startswith('`') and part.endswith('`'):
                    run = p.add_run(part[1:-1])
                    run.font.name = 'Consolas'
                    run.font.size = Pt(9)
                else:
                    p.add_run(part)

        i += 1

    doc.save(DOCX_OUT)
    print(f"Generated concise DOCX report: {DOCX_OUT}")

if __name__ == "__main__":
    build_docx()
